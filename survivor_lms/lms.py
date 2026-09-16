"""Real "Last Man Standing" pool data: everyone's picks + real pick-count %,
parsed straight from the pool operator's exported workbook (no scraping/heuristics
needed once this file is dropped in).

Files, all under data/<season>/lms/ :

  *.xlsx                     the pool operator's export (e.g. "WEEKLY PICKS.xlsx").
                              Any sheet named "WEEK <n>" is read as that week's
                              picks: column A = entry name, column B = team picked.
                              Any sheet starting with "COUNT" is read as the real
                              pick-count pivot (Row Labels = team, next col = count,
                              a "Grand Total" row = field size).
  members.csv                member (your nickname) -> entry_name (exact string as
                              it appears in the workbook). One row per person in
                              your group you want a personalised recommendation for.
  used_teams.csv             legacy single-user fallback, only consulted when no
                              --member is given or the member has no workbook history.
  participants.csv           optional hand-logged side-pool opponent picks (same
                              schema as constraints.Roster), independent of members.csv.
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

from .config import Config
from survivor_core.teams import ABBRS, to_abbr

WEEK_SHEET_RE = re.compile(r"^\s*WEEK\s*(\d+)\s*$", re.IGNORECASE)
COUNT_SHEET_RE = re.compile(r"^\s*COUNT", re.IGNORECASE)


def lms_dir(cfg: Config) -> Path:
    d = cfg.year_dir / "lms"
    d.mkdir(parents=True, exist_ok=True)
    return d


def find_workbook(cfg: Config) -> Path | None:
    """Most-recently-modified *.xlsx in lms/, ignoring Excel's ~$lock files."""
    candidates = [p for p in lms_dir(cfg).glob("*.xlsx") if not p.name.startswith("~$")]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def load_members(cfg: Config) -> dict[str, str]:
    """member (short name you call them) -> pool entry_name, from lms/members.csv."""
    f = lms_dir(cfg) / "members.csv"
    if not f.exists():
        f.write_text(
            "member,entry_name\n"
            "Mattymo,Mattymo\n"
            "# member = short name you refer to them by; entry_name = EXACT name/handle\n"
            "# as it appears in the pool operator's export (case-sensitive).\n"
            "# Add one row per person in your group you want a recommendation for.\n"
        )
    df = pd.read_csv(f, dtype=str, comment=None).fillna("")
    df = df[df["member"].str.strip().ne("") & ~df["member"].str.strip().str.startswith("#")]
    return {r["member"].strip(): r["entry_name"].strip() for _, r in df.iterrows()}


def load_all_entries(cfg: Config) -> pd.DataFrame:
    """Every entrant's pick in every 'WEEK n' sheet -> long DataFrame(entry, week, team).

    Unparsable team names (typos, blanks, mid-season header rows) are silently
    skipped rather than raising -- this file has ~14k rows of messy human input.
    """
    wb_path = find_workbook(cfg)
    if wb_path is None:
        return pd.DataFrame(columns=["entry", "week", "team"])
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    rows = []
    for name in wb.sheetnames:
        m = WEEK_SHEET_RE.match(name)
        if not m:
            continue
        week = int(m.group(1))
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        next(it, None)  # header row
        for r in it:
            if not r or r[0] is None or len(r) < 2 or r[1] is None:
                continue
            entry = str(r[0]).strip()
            if not entry:
                continue
            try:
                team = to_abbr(str(r[1]).strip())
            except KeyError:
                continue
            rows.append((entry, week, team))
    return pd.DataFrame(rows, columns=["entry", "week", "team"])


def load_pick_pct(cfg: Config, week: int) -> tuple[pd.Series | None, dict]:
    """Real public pick % for `week` from the "COUNT OF SELECTED TEAMS" pivot sheet.

    Returns (Series abbr -> fraction over the 32 teams, meta). Returns (None, meta)
    if there's no workbook, no count sheet, or the count sheet is for a different
    week than asked (the pool operator only publishes counts for the live week).
    """
    wb_path = find_workbook(cfg)
    if wb_path is None:
        return None, {"source": "lms-count", "note": "no workbook found in lms/"}
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    sheet = next((n for n in wb.sheetnames if COUNT_SHEET_RE.match(n)), None)
    if sheet is None:
        return None, {"source": "lms-count", "note": "no COUNT... sheet found"}
    ws = wb[sheet]
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None]

    counts_week = None
    for r in rows[:5]:
        if len(r) > 1 and r[1]:
            mm = re.search(r"WEEK\s*(\d+)", str(r[1]), re.IGNORECASE)
            if mm:
                counts_week = int(mm.group(1))
                break
    if counts_week is not None and counts_week != week:
        return None, {"source": "lms-count", "sheet": sheet, "week": counts_week,
                      "note": f"count sheet is for Week {counts_week}, asked for Week {week}"}

    counts: dict[str, float] = {}
    total = 0.0
    for r in rows:
        label = str(r[0]).strip()
        val = r[1] if len(r) > 1 else None
        if not label or val is None:
            continue
        low = label.lower()
        if low in ("row labels",):
            continue
        if low == "grand total":
            total = float(val)
            continue
        try:
            ab = to_abbr(label)
        except KeyError:
            continue
        counts[ab] = float(val)

    if not counts:
        return None, {"source": "lms-count", "sheet": sheet, "note": "could not parse any team rows"}
    if not total:
        total = sum(counts.values())

    full = pd.Series(0.0, index=ABBRS)
    full.update({k: v / total for k, v in counts.items()})
    return full, {"source": "lms-count", "sheet": sheet, "week": counts_week or week, "n_entries": int(total)}


def member_used_teams(entries: pd.DataFrame, entry_name: str, before_week: int) -> set[str]:
    if entries.empty:
        return set()
    m = entries["entry"].eq(entry_name) & entries["week"].between(1, before_week - 1)
    return set(entries.loc[m, "team"])


def member_known_pick(entries: pd.DataFrame, entry_name: str, week: int) -> str | None:
    if entries.empty:
        return None
    m = entries["entry"].eq(entry_name) & entries["week"].eq(week)
    vals = entries.loc[m, "team"].tolist()
    return vals[0] if vals else None
